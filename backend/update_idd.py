"""Update DATS_Simulator_IDD_20261010.docx to align with the MuST / DATS RF Simulator project.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Century Gothic"
COLOR_HEADER_BG = "F1A983"      # Primary header fill (peach/coral)
COLOR_WORD_HEADER_BG = "FAE2D5" # Word W00..W15 header fill (light peach)
COLOR_META_BG = "DFEDEA"        # Frame metadata fill (light cyan/teal)
COLOR_PAYLOAD_BG = "B3E5A1"     # Payload data fill (soft green)
COLOR_WHITE = "FFFFFF"

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def format_cell(cell, text, font_size=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, fill_hex=None):
    if fill_hex:
        set_cell_shading(cell, fill_hex)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(font_size)
    run.font.bold = bold

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_p(doc, text, level=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(13 if level == 2 else 11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64) if level == 2 else RGBColor(0x33, 0x33, 0x33)
    return p

def add_caption_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(9.5)
    run.font.italic = True
    run.font.bold = True
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    return p

def create_word_layout_table(doc, rows_data, caption_text):
    """Creates a 16-bit word layout table (W00-W15 + BYTE offset).
    rows_data is a list of tuples: (left_text, right_text, byte_offset, fill_hex)
    If right_text is None, left_text spans the full 16-bit word.
    """
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=17)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    # Header row: W00..W15 + BYTE
    for col_idx in range(16):
        format_cell(tbl.cell(0, col_idx), f"W{col_idx:02d}", font_size=7.5, bold=True, fill_hex=COLOR_WORD_HEADER_BG)
    format_cell(tbl.cell(0, 16), "BYTE", font_size=7.5, bold=True, fill_hex=COLOR_WORD_HEADER_BG)
    
    for r_idx, (left_txt, right_txt, byte_offset, fill) in enumerate(rows_data, start=1):
        if right_txt is None:
            # Span across all 16 bit columns
            tbl.cell(r_idx, 0).merge(tbl.cell(r_idx, 15))
            format_cell(tbl.cell(r_idx, 0), left_txt, font_size=9.0, bold=False, fill_hex=fill)
        else:
            # W00..W07 (cols 0..7)
            tbl.cell(r_idx, 0).merge(tbl.cell(r_idx, 7))
            format_cell(tbl.cell(r_idx, 0), left_txt, font_size=9.0, bold=False, fill_hex=fill)
            # W08..W15 (cols 8..15)
            tbl.cell(r_idx, 8).merge(tbl.cell(r_idx, 15))
            format_cell(tbl.cell(r_idx, 8), right_txt, font_size=9.0, bold=False, fill_hex=fill)
        format_cell(tbl.cell(r_idx, 16), str(byte_offset), font_size=9.0, bold=True, fill_hex=COLOR_WHITE)
    
    add_caption_p(doc, caption_text)
    return tbl

def create_field_desc_table(doc, headers, rows_data, caption_text, col_widths=None):
    """Creates a standard field description table."""
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    # Headers
    for c_idx, h in enumerate(headers):
        format_cell(tbl.cell(0, c_idx), h, font_size=9.5, bold=True, fill_hex=COLOR_HEADER_BG)
    
    # Rows
    for r_idx, row in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(row):
            # Left align notes/descriptions (last column)
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == len(row) - 1 or c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(tbl.cell(r_idx, c_idx), str(val), font_size=9.0, bold=False, align=align, fill_hex=COLOR_WHITE)
            
    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    add_caption_p(doc, caption_text)
    return tbl

def main():
    doc = docx.Document("DATS_Simulator_IDD_20261010.docx")
    
    print("1. Updating Cover / Title...")
    for p in doc.paragraphs:
        if p.text.strip() == "DATS":
            p.text = ""
            r = p.add_run("DATS\nMulti-Carrier Satellite Telemetry (MuST)\nRF Simulator")
            r.font.name = FONT_NAME
            r.font.size = Pt(28)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    print("2. Updating Table 1 (List of Messages)...")
    t1 = doc.tables[0]
    # Keep header row, clear/add rows
    while len(t1.rows) > 1:
        t1._tbl.remove(t1.rows[-1]._tr)
        
    t1_headers = ["S.no", "Message Type", "Source", "Destination", "Msg ID", "Trigger", "Section"]
    for c_idx, h in enumerate(t1_headers):
        format_cell(t1.cell(0, c_idx), h, font_size=9.5, bold=True, fill_hex=COLOR_HEADER_BG)
        
    t1_data = [
        ("1", "RF Control", "Console (0x01)", "Simulator (0x02)", "0x01", "Operator Start / Stop RF button", "4.2"),
        ("2", "Set Mode", "Console (0x01)", "Simulator (0x02)", "0x02", "Operator Mode switch (Simulation / Radiation)", "4.3"),
        ("3", "Parameter Load / Carrier Configuration", "Console (0x01)", "Simulator (0x02)", "0x10", "Operator Apply Configuration (Carriers 1 to 12)", "4.1"),
        ("4", "IQ Stream Begin", "Console (0x01)", "Simulator (0x02)", "0x21", "Operator 'Send to Carrier N' action in Simulation", "4.4"),
        ("5", "IQ Stream Data", "Console (0x01)", "Simulator (0x02)", "0x20", "Active arbitrary baseband I/Q sample stream chunk", "4.5"),
        ("6", "IQ Stream End", "Console (0x01)", "Simulator (0x02)", "0x22", "Completion of arbitrary baseband I/Q sample stream", "4.6"),
        ("7", "Command Acknowledgment (ACK)", "Simulator (0x02)", "Console (0x01)", "Matches Cmd (Kind 0xBB)", "Positive ACK / Execution status response from hardware", "4.7"),
    ]
    for row in t1_data:
        r = t1.add_row()
        for c_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx in (1, 5) else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(r.cells[c_idx], val, font_size=9.0, bold=False, align=align, fill_hex=COLOR_WHITE)

    # Update caption for Table 1
    for p in doc.paragraphs:
        if "Table" in p.text and "List of Messages" in p.text:
            p.text = "Table 1: List of Messages"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    print("3. Updating Table 2 (Frame Structure Caption)...")
    for p in doc.paragraphs:
        if "Table" in p.text and "Frame Structure" in p.text:
            p.text = "Table 2: Common Frame Structure (29 + N Bytes)"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    print("4. Updating Table 3 (Description of Common Message Fields)...")
    t3 = doc.tables[2]
    while len(t3.rows) > 1:
        t3._tbl.remove(t3.rows[-1]._tr)
        
    t3_headers = ["Name of the field", "Size (bytes)", "Notes"]
    for c_idx, h in enumerate(t3_headers):
        format_cell(t3.cell(0, c_idx), h, font_size=9.5, bold=True, fill_hex=COLOR_HEADER_BG)
        
    t3_data = [
        ("SOP", "4", "Start of Packet: Fixed 4-byte synchronization magic word 0xB1B1B1B1 (0xB1, 0xB1, 0xB1, 0xB1)."),
        ("Destination ID", "1", "CSCI ID of destination subsystem (0x02 = MuST/DATS RF Simulator, 0x01 = Console)."),
        ("Source ID", "1", "CSCI ID of source subsystem (0x01 = Operator Console, 0x02 = Simulator)."),
        ("Destination Subsystem ID", "1", "Destination subsystem identifier (default 0x00)."),
        ("Source Subsystem ID", "1", "Source subsystem identifier (default 0x00)."),
        ("Destination Sub-subsystem ID", "1", "Destination sub-subsystem identifier (default 0x00)."),
        ("Source Sub-subsystem ID", "1", "Source sub-subsystem identifier (default 0x00)."),
        ("Message Counter", "1", "Incremental frame sequence counter (0 to 255), incremented per transmitted frame."),
        ("Command ID", "1", "Identifies the command type as defined in Table 1 (e.g. 0x01, 0x02, 0x10, 0x20, 0x21, 0x22)."),
        ("Reserved", "1", "Reserved byte for future protocol expansion (set to 0x00)."),
        ("Data / Ack (Packet Kind)", "1", "Packet classification flag: 0xAA = Command / Data packet, 0xBB = Acknowledgment (ACK)."),
        ("Timestamp", "8", "64-bit integer representing elapsed seconds since Unix epoch (00:00:00 UTC on 1 January 1970), Big-Endian."),
        ("Data Length (N)", "2", "16-bit unsigned integer representing length of the Data Packet payload in bytes (Big-Endian)."),
        ("Data Packet (Payload)", "N", "Command-specific payload data (length N bytes as specified in the Data Length field)."),
        ("Checksum", "1", "8-bit XOR checksum computed across all bytes from SOP (byte 0) through end of Data Packet (byte 23+N)."),
        ("EOP", "4", "End of Packet: Fixed 4-byte termination magic word 0xE1E1E1E1 (0xE1, 0xE1, 0xE1, 0xE1)."),
    ]
    for row in t3_data:
        r = t3.add_row()
        for c_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 2 or c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(r.cells[c_idx], val, font_size=9.0, bold=False, align=align, fill_hex=COLOR_WHITE)

    for p in doc.paragraphs:
        if "Table" in p.text and "Description of the Message" in p.text:
            p.text = "Table 3: Field Description of Common Frame Structure"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    print("5. Updating Table 4 (Port, Protocol and Network Configuration)...")
    t4 = doc.tables[3]
    while len(t4.rows) > 1:
        t4._tbl.remove(t4.rows[-1]._tr)
        
    t4_headers = ["Interface", "Protocol", "IP Address", "Port"]
    for c_idx, h in enumerate(t4_headers):
        format_cell(t4.cell(0, c_idx), h, font_size=9.5, bold=True, fill_hex=COLOR_HEADER_BG)
        
    t4_data = [
        ("RFSoC Simulator Hardware (PS GEM Ethernet)", "TCP (lwIP server)", "192.168.1.10 (Default static IP)", "7000"),
        ("RFSoC Simulator Hardware (Alternative Link)", "UDP (lwIP server)", "192.168.1.10 (Default static IP)", "7000"),
        ("Offline Simulator / Mock Hardware Device", "TCP / UDP", "127.0.0.1 (Local loopback)", "7000"),
        ("Operator Workstation Host PC NIC", "TCP / UDP Client", "192.168.1.x / 24 subnet (e.g. 192.168.1.100)", "Ephemeral / Dynamic"),
        ("Operator Web Console UI & REST API Server", "HTTP / WebSocket", "127.0.0.1 / localhost", "8765"),
    ]
    for row in t4_data:
        r = t4.add_row()
        for c_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx in (0, 2) else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(r.cells[c_idx], val, font_size=9.0, bold=False, align=align, fill_hex=COLOR_WHITE)

    for p in doc.paragraphs:
        if "Table" in p.text and "Port" in p.text and "Network" in p.text:
            p.text = "Table 4: Port, Protocol and Network Configuration"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    print("6. Updating Table 5 and Table 6 (Parameter Load / Carrier Configuration)...")
    # Update Section 4 headings
    for p in doc.paragraphs:
        if p.text.strip() == "Parameter Load":
            p.text = "4.1 Parameter Load / Carrier Configuration Command (Msg ID: 0x10)"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(12)
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    # Table 5: 16-bit word layout for CARRIER_CONFIG (17 bytes)
    t5 = doc.tables[4]
    while len(t5.rows) > 1:
        t5._tbl.remove(t5.rows[-1]._tr)
        
    for col_idx in range(16):
        format_cell(t5.cell(0, col_idx), f"W{col_idx:02d}", font_size=7.5, bold=True, fill_hex=COLOR_WORD_HEADER_BG)
    format_cell(t5.cell(0, 16), "BYTE", font_size=7.5, bold=True, fill_hex=COLOR_WORD_HEADER_BG)

    t5_word_rows = [
        ("Carrier Number (uint8: 1 to 12)", "Carrier Enable (uint8: 0=Off, 1=On)", "0", COLOR_PAYLOAD_BG),
        ("Waveform Code (uint8: 0 to 4)", "Bit Rate kbps [MSB] (uint16)", "2", COLOR_PAYLOAD_BG),
        ("Bit Rate kbps [LSB]", "Subcarrier kHz [MSB] (uint16)", "4", COLOR_PAYLOAD_BG),
        ("Subcarrier kHz [LSB]", "PCM Code (uint8: 0=NRZ-L, 1=NRZ-S)", "6", COLOR_PAYLOAD_BG),
        ("Convolutional Coding (uint8: 0=Off, 1=On)", "Frame Length bytes [MSB] (uint16)", "8", COLOR_PAYLOAD_BG),
        ("Frame Length bytes [LSB]", "Frequency Offset Hz [Byte 3 / MSB] (int32)", "10", COLOR_PAYLOAD_BG),
        ("Frequency Offset Hz [Byte 2, Byte 1] (int32)", None, "12", COLOR_PAYLOAD_BG),
        ("Frequency Offset Hz [Byte 0 / LSB]", "Output Power dBm [MSB] (int16)", "14", COLOR_PAYLOAD_BG),
        ("Output Power dBm [LSB]", "Reserved / Padding (0x00)", "16", COLOR_PAYLOAD_BG),
    ]
    for left_txt, right_txt, byte_offset, fill in t5_word_rows:
        r = t5.add_row()
        r_idx = len(t5.rows) - 1
        if right_txt is None:
            t5.cell(r_idx, 0).merge(t5.cell(r_idx, 15))
            format_cell(t5.cell(r_idx, 0), left_txt, font_size=8.5, bold=False, fill_hex=fill)
        else:
            t5.cell(r_idx, 0).merge(t5.cell(r_idx, 7))
            format_cell(t5.cell(r_idx, 0), left_txt, font_size=8.5, bold=False, fill_hex=fill)
            t5.cell(r_idx, 8).merge(t5.cell(r_idx, 15))
            format_cell(t5.cell(r_idx, 8), right_txt, font_size=8.5, bold=False, fill_hex=fill)
        format_cell(t5.cell(r_idx, 16), byte_offset, font_size=8.5, bold=True, fill_hex=COLOR_WHITE)

    for p in doc.paragraphs:
        if "Table" in p.text and "Beam Schedule Command" in p.text and "Field" not in p.text:
            p.text = "Table 5: Frame Layout of Parameter Load / Carrier Configuration Command (17 Bytes)"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Table 6: Field Description of Carrier Configuration Command
    t6 = doc.tables[5]
    while len(t6.rows) > 1:
        t6._tbl.remove(t6.rows[-1]._tr)
        
    t6_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    for c_idx, h in enumerate(t6_headers):
        format_cell(t6.cell(0, c_idx), h, font_size=9.5, bold=True, fill_hex=COLOR_HEADER_BG)

    t6_data = [
        ("Carrier Number", "1", "1 to 12", "Target carrier channel index to configure."),
        ("Carrier Enable", "1", "0: Disabled, 1: Enabled", "Switches carrier RF transmission on or off."),
        ("Waveform / Modulation", "1", "0: PCM/FSK\n1: PCM/PSK\n2: PCM/PSK/PM\n3: CDMA\n4: IQ Data", "Selects the active modulation waveform type."),
        ("Bit Rate", "2", "4, 16, 32, 64, 128, 256, 512, 1024, 2048", "Data transmission rate in kbps (Big-Endian uint16). Used for PCM/FSK, PCM/PSK, PCM/PSK/PM, CDMA."),
        ("Subcarrier Frequency", "2", "32 to 256", "Subcarrier frequency in kHz (Big-Endian uint16). Required for PCM/PSK/PM; set to 0 when not applicable."),
        ("PCM Code", "1", "0: NRZ-L\n1: NRZ-S", "Pulse Code Modulation line code format."),
        ("Convolutional Coding", "1", "0: Disabled, 1: Enabled", "Forward Error Correction: Rate 1/2, constraint length k=7 (standard for satellite telemetry)."),
        ("Frame Length", "2", "128 to 256", "Telemetry data frame length in bytes (Big-Endian uint16)."),
        ("Frequency Offset", "4", "-50,000,000 to +50,000,000", "Carrier center frequency offset in Hz (Big-Endian signed 32-bit int), relative to nominal 2.25 GHz center."),
        ("Output Power", "2", "-150 to +20", "Carrier output power in dBm (Big-Endian signed 16-bit int) in 1 dB steps."),
    ]
    for row in t6_data:
        r = t6.add_row()
        for c_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx in (0, 3) else WD_ALIGN_PARAGRAPH.CENTER
            format_cell(r.cells[c_idx], val, font_size=9.0, bold=False, align=align, fill_hex=COLOR_WHITE)

    for p in doc.paragraphs:
        if "Table" in p.text and "Field Description of Beam Schedule Command" in p.text:
            p.text = "Table 6: Field Description of Parameter Load / Carrier Configuration Command"
            p.runs[0].font.name = FONT_NAME
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.bold = True
            p.runs[0].font.italic = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    print("7. Adding additional Command specifications to Section 4...")
    # 4.2 RF Control Command
    add_heading_p(doc, "4.2 RF Control Command (Msg ID: 0x01)", level=2)
    add_body_p(doc, "The RF Control command starts or stops the RFSoC Digital-to-Analog Converter (DAC) output transmissions. In Radiation Mode, the host operator interface requires an explicit safety confirmation prior to initiating RF radiation.")
    
    rf_word_rows = [
        ("RF Output State (uint8: 0=Stop RF, 1=Start RF)", "Reserved / Padding (0x00)", "0", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, rf_word_rows, "Table 7: Frame Layout of RF Control Command (1 Byte Payload)")
    
    rf_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    rf_desc_rows = [
        ("RF Output State", "1", "0x00: Stop RF\n0x01: Start RF", "0x00 shuts down/mutes RF DAC outputs; 0x01 enables DAC transmission."),
    ]
    create_field_desc_table(doc, rf_desc_headers, rf_desc_rows, "Table 8: Field Description of RF Control Command")

    # 4.3 Mode Selection Command
    add_heading_p(doc, "4.3 Mode Selection Command (Msg ID: 0x02)", level=2)
    add_body_p(doc, "The Mode Selection command toggles the simulator between Simulation Mode and Radiation Mode. In Simulation Mode, the operator can load captured waveforms (e.g. Vivado ILA CSV or I/Q files), inspect time-domain signals, spectrum, and constellations, and stream samples to any carrier. In Radiation Mode, DACs transmit over-the-air with safety interlocks.")
    
    mode_word_rows = [
        ("Operating Mode (uint8: 0=Simulation, 1=Radiation)", "Reserved / Padding (0x00)", "0", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, mode_word_rows, "Table 9: Frame Layout of Mode Selection Command (1 Byte Payload)")
    
    mode_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    mode_desc_rows = [
        ("Operating Mode", "1", "0x00: Simulation Mode\n0x01: Radiation Mode", "0x00 enables offline/streamed waveform simulation; 0x01 enables over-the-air RF radiation mode."),
    ]
    create_field_desc_table(doc, mode_desc_headers, mode_desc_rows, "Table 10: Field Description of Mode Selection Command")

    # 4.4 IQ Stream Begin Command
    add_heading_p(doc, "4.4 IQ Stream Begin Command (Msg ID: 0x21)", level=2)
    add_body_p(doc, "The IQ Stream Begin command prepares the target carrier on the RFSoC simulator to receive an arbitrary baseband sample stream. It declares the baseband sample rate and total sample count expected.")
    
    iq_begin_word_rows = [
        ("Carrier Number (uint8: 1 to 12)", "Sample Rate Hz [Byte 3 / MSB] (uint32)", "0", COLOR_PAYLOAD_BG),
        ("Sample Rate Hz [Byte 2, Byte 1] (uint32)", None, "2", COLOR_PAYLOAD_BG),
        ("Sample Rate Hz [Byte 0 / LSB]", "Total Samples [Byte 3 / MSB] (uint32)", "4", COLOR_PAYLOAD_BG),
        ("Total Samples [Byte 2, Byte 1] (uint32)", None, "6", COLOR_PAYLOAD_BG),
        ("Total Samples [Byte 0 / LSB]", "Reserved / Padding (0x00)", "8", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, iq_begin_word_rows, "Table 11: Frame Layout of IQ Stream Begin Command (9 Bytes Payload)")
    
    iq_begin_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    iq_begin_desc_rows = [
        ("Carrier Number", "1", "1 to 12", "Target carrier channel on which baseband I/Q samples will be played."),
        ("Sample Rate", "4", "1 to 100,000,000", "Baseband sampling rate in Hz (Big-Endian uint32)."),
        ("Total Samples", "4", "1 to 4,294,967,295", "Total number of complex (I, Q) sample pairs in the upcoming stream (Big-Endian uint32)."),
    ]
    create_field_desc_table(doc, iq_begin_desc_headers, iq_begin_desc_rows, "Table 12: Field Description of IQ Stream Begin Command")

    # 4.5 IQ Stream Data Command
    add_heading_p(doc, "4.5 IQ Stream Data Command (Msg ID: 0x20)", level=2)
    add_body_p(doc, "The IQ Stream Data command transfers a chunk of interleaved 16-bit signed I and Q baseband samples to the target carrier. Chunks are numbered sequentially, and each packet contains up to 1024 sample pairs.")
    
    iq_data_word_rows = [
        ("Carrier Number (uint8: 1 to 12)", "Chunk Sequence Number [Byte 3 / MSB] (uint32)", "0", COLOR_PAYLOAD_BG),
        ("Chunk Sequence Number [Byte 2, Byte 1] (uint32)", None, "2", COLOR_PAYLOAD_BG),
        ("Chunk Sequence Number [Byte 0 / LSB]", "Sample Count N [MSB] (uint16)", "4", COLOR_PAYLOAD_BG),
        ("Sample Count N [LSB]", "I[0] Sample [MSB] (int16)", "6", COLOR_PAYLOAD_BG),
        ("I[0] Sample [LSB]", "Q[0] Sample [MSB] (int16)", "8", COLOR_PAYLOAD_BG),
        ("...", "...", "...", COLOR_PAYLOAD_BG),
        ("I[N-1] Sample (int16)", "Q[N-1] Sample (int16)", "7 + 4*(N-1)", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, iq_data_word_rows, "Table 13: Frame Layout of IQ Stream Data Command (7 + 4*N Bytes Payload)")
    
    iq_data_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    iq_data_desc_rows = [
        ("Carrier Number", "1", "1 to 12", "Target carrier channel."),
        ("Chunk Sequence Number", "4", "0, 1, 2, ...", "Zero-based sequential chunk counter (Big-Endian uint32)."),
        ("Sample Count (N)", "2", "1 to 1024", "Number of complex (I, Q) sample pairs contained in this packet (Big-Endian uint16)."),
        ("Interleaved I/Q Samples", "4*N", "-32768 to +32767", "N pairs of 16-bit signed integers in big-endian byte order: I0, Q0, I1, Q1, ..., I[N-1], Q[N-1]."),
    ]
    create_field_desc_table(doc, iq_data_desc_headers, iq_data_desc_rows, "Table 14: Field Description of IQ Stream Data Command")

    # 4.6 IQ Stream End Command
    add_heading_p(doc, "4.6 IQ Stream End Command (Msg ID: 0x22)", level=2)
    add_body_p(doc, "The IQ Stream End command notifies the RFSoC simulator firmware that all sample chunks for the specified carrier have been transmitted, indicating the total number of chunks sent.")
    
    iq_end_word_rows = [
        ("Carrier Number (uint8: 1 to 12)", "Total Chunks [Byte 3 / MSB] (uint32)", "0", COLOR_PAYLOAD_BG),
        ("Total Chunks [Byte 2, Byte 1] (uint32)", None, "2", COLOR_PAYLOAD_BG),
        ("Total Chunks [Byte 0 / LSB]", "Reserved / Padding (0x00)", "4", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, iq_end_word_rows, "Table 15: Frame Layout of IQ Stream End Command (5 Bytes Payload)")
    
    iq_end_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    iq_end_desc_rows = [
        ("Carrier Number", "1", "1 to 12", "Target carrier channel."),
        ("Total Chunks", "4", "1 to 4,294,967,295", "Total number of IQ_DATA chunks successfully transmitted (Big-Endian uint32)."),
    ]
    create_field_desc_table(doc, iq_end_desc_headers, iq_end_desc_rows, "Table 16: Field Description of IQ Stream End Command")

    # 4.7 Acknowledge Frame (ACK)
    add_heading_p(doc, "4.7 Command Acknowledgment Frame (ACK, Packet Kind: 0xBB)", level=2)
    add_body_p(doc, "The simulator answers every valid received command with an Acknowledge (ACK) frame within the configured ACK timeout (default: 2.0 s). In the ACK frame header, Destination ID and Source ID are reversed (Destination=0x01, Source=0x02), Message Counter matches the command sequence counter, Packet Kind is set to 0xBB, and Command ID matches the received command. The single-byte payload reports the execution status code.")
    
    ack_word_rows = [
        ("Execution Status Code (uint8: 0=OK, Non-zero=Error)", "Reserved / Padding (0x00)", "0", COLOR_PAYLOAD_BG),
    ]
    create_word_layout_table(doc, ack_word_rows, "Table 17: Frame Layout of Acknowledge Frame (1 Byte Payload)")
    
    ack_desc_headers = ["Name of the field", "Size (byte)", "Value / Range", "Notes"]
    ack_desc_rows = [
        ("Execution Status Code", "1", "0x00: Success / OK\n0x01: Invalid Parameter\n0x02: Checksum Mismatch\n0x03: State Error", "0x00 indicates valid frame receipt and successful hardware execution; non-zero values denote error rejection."),
    ]
    create_field_desc_table(doc, ack_desc_headers, ack_desc_rows, "Table 18: Field Description of Acknowledge Frame")

    doc.save("DATS_Simulator_IDD_20261010.docx")
    print("SUCCESS: DATS_Simulator_IDD_20261010.docx updated successfully!")

if __name__ == "__main__":
    main()
